"""Parse the workbook timetable without guessing unresolved academic mappings."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


CLASS_COLUMNS = {6: "D12A", 7: "D12B", 8: "D12C"}
CANONICAL_SUBJECTS = {"SSD", "CCL", "DE", "ML", "CSS", "DWM", "SEA", "TCS"}
DAY_PATTERN = re.compile(r"^(MONDAY|TUESDAY|WEDNESDAY|THURSDAY|FRIDAY)\b", re.IGNORECASE)
TIME_PATTERN = re.compile(r"^(\d{1,2})[.:](\d{2})\s*-\s*(\d{1,2})[.:](\d{2})$")
COMBINED_PATTERN = re.compile(
    r"^(?P<subject>[A-Z0-9]+)\s*/\s*\((?P<faculty>[A-Z0-9]+)\)\s*/\s*"
    r"(?P<room>[A-Z0-9]+)\s*\((?P<classes>[^)]+)\)$",
    re.IGNORECASE,
)
PRACTICAL_PATTERN = re.compile(
    r"^(?P<subject>[A-Z0-9]+)\s*/\s*(?P<batch>[A-Z0-9]+)\s*/\s*"
    r"(?P<room>[A-Z0-9]+)\s*/\s*(?P<faculty>[A-Z0-9]+)$",
    re.IGNORECASE,
)
PROJECT_PATTERN = re.compile(r"^(?:(?P<batch>[A-Z0-9]+)\s*/\s*)?PROJECT$", re.IGNORECASE)
SIMPLE_PATTERN = re.compile(
    r"^(?P<subject>[A-Z0-9]+)\s+\((?P<faculty>[A-Z0-9]+)\)\s+(?P<room>[A-Z0-9]+)$",
    re.IGNORECASE,
)


@dataclass
class ParsedOccurrence:
    day: str
    start_time: str
    end_time: str
    source_cell: str
    source_item_index: int
    raw_source_text: str
    source_subject_code: str | None = None
    course_id: str | None = None
    faculty_initial: str | None = None
    faculty_id: str | None = None
    classroom: str | None = None
    class_ids: list[str] = field(default_factory=list)
    batch_codes: list[str] = field(default_factory=list)
    activity_kind: str = "lecture"
    parallel_group: str | None = None
    issues: list[str] = field(default_factory=list)


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def parse_time(value: Any) -> tuple[str, str] | None:
    match = TIME_PATTERN.fullmatch(clean(value))
    if not match:
        return None
    start_hour, start_minute, end_hour, end_minute = (int(part) for part in match.groups())
    if not (0 <= start_minute < 60 and 0 <= end_minute < 60):
        return None

    # The supplied timetable omits AM/PM. Its teaching day begins at 08:30,
    # so 01:30, 02:30 and 03:30 are afternoon values. Preserve explicit
    # 24-hour values while normalizing the ambiguous source format.
    if 1 <= start_hour <= 3:
        start_hour += 12
    start_total = start_hour * 60 + start_minute
    end_total = end_hour * 60 + end_minute
    if end_total <= start_total:
        end_hour += 12
    if start_hour > 23 or end_hour > 23:
        return None
    return f"{start_hour:02d}:{start_minute:02d}", f"{end_hour:02d}:{end_minute:02d}"


def cell_time_range(worksheet: Any, row_number: int, column: int) -> tuple[str, str] | None:
    """Return the full duration represented by a timetable cell.

    Excel uses vertical merges for two-hour practicals. Only the top-left
    cell carries the value, so its end time must come from the final merged
    row rather than from the first one-hour slot.
    """
    base = parse_time(worksheet.cell(row_number, 5).value)
    if not base:
        return None
    for merged in worksheet.merged_cells.ranges:
        if (
            merged.min_row == row_number
            and merged.max_row > row_number
            and merged.min_col <= column <= merged.max_col
        ):
            final_slot = parse_time(worksheet.cell(merged.max_row, 5).value)
            if final_slot:
                return base[0], final_slot[1]
    return base


def load_mappings(path: Path | None) -> dict[str, Any]:
    if not path:
        return {"subject_aliases": {}, "class_aliases": {}, "project_activity_approved": False}
    mappings = json.loads(path.read_text(encoding="utf-8"))
    return {
        "subject_aliases": mappings.get("subject_aliases", {}),
        "class_aliases": mappings.get("class_aliases", {}),
        "project_activity_approved": bool(mappings.get("project_activity_approved", False)),
    }


def faculty_initials(workbook: Any) -> set[str]:
    worksheet = workbook["Faculty"]
    headers = [clean(worksheet.cell(1, column).value).casefold() for column in range(1, worksheet.max_column + 1)]
    initial_column = headers.index("initial") + 1
    return {
        clean(worksheet.cell(row, initial_column).value).upper()
        for row in range(2, worksheet.max_row + 1)
        if clean(worksheet.cell(row, initial_column).value)
    }


def split_items(raw: str) -> list[str]:
    blocks = [re.sub(r"\s+", " ", block).strip() for block in re.split(r"\n\s*\n", raw) if block.strip()]
    if len(blocks) == 1:
        lines = [line.strip() for line in raw.splitlines() if line.strip()]
        if len(lines) > 1 and all(COMBINED_PATTERN.fullmatch(line) for line in lines):
            return lines
    return blocks


def resolve_subject(token: str, mappings: dict[str, Any], occurrence: ParsedOccurrence) -> None:
    normalized = token.upper()
    mapped = mappings["subject_aliases"].get(normalized, normalized)
    occurrence.source_subject_code = normalized
    if not mapped or mapped not in CANONICAL_SUBJECTS:
        occurrence.issues.append(f"unresolved subject token {normalized}")
    else:
        occurrence.course_id = mapped


def resolve_faculty(initial: str, known_initials: set[str], occurrence: ParsedOccurrence) -> None:
    normalized = initial.upper()
    occurrence.faculty_initial = normalized
    if normalized not in known_initials:
        occurrence.issues.append(f"unknown faculty initial {normalized}")
    else:
        occurrence.faculty_id = f"FAC-{re.sub(r'[^A-Z0-9]', '', normalized)}"


def resolve_classes(tokens: list[str], mappings: dict[str, Any], occurrence: ParsedOccurrence) -> None:
    for token in tokens:
        normalized = token.strip().upper()
        mapped = mappings["class_aliases"].get(normalized, normalized)
        if not mapped or mapped not in CLASS_COLUMNS.values():
            occurrence.issues.append(f"unresolved class token {normalized}")
        elif mapped not in occurrence.class_ids:
            occurrence.class_ids.append(mapped)


def parse_item(
    item: str,
    *,
    day: str,
    time_range: tuple[str, str],
    coordinate: str,
    item_index: int,
    default_class: str,
    item_count: int,
    mappings: dict[str, Any],
    known_initials: set[str],
) -> ParsedOccurrence:
    occurrence = ParsedOccurrence(
        day=day,
        start_time=time_range[0],
        end_time=time_range[1],
        source_cell=coordinate,
        source_item_index=item_index,
        raw_source_text=item,
        parallel_group=coordinate if item_count > 1 else None,
    )
    combined = COMBINED_PATTERN.fullmatch(item)
    practical = PRACTICAL_PATTERN.fullmatch(item)
    project = PROJECT_PATTERN.fullmatch(item)
    simple = SIMPLE_PATTERN.fullmatch(item)

    if combined:
        resolve_subject(combined["subject"], mappings, occurrence)
        resolve_faculty(combined["faculty"], known_initials, occurrence)
        occurrence.classroom = combined["room"].upper()
        resolve_classes(combined["classes"].split("+"), mappings, occurrence)
    elif practical:
        resolve_subject(practical["subject"], mappings, occurrence)
        resolve_faculty(practical["faculty"], known_initials, occurrence)
        occurrence.classroom = practical["room"].upper()
        occurrence.class_ids = [default_class]
        occurrence.batch_codes = [practical["batch"].upper()]
        occurrence.activity_kind = "practical"
    elif project:
        occurrence.source_subject_code = "PROJECT"
        occurrence.class_ids = [default_class]
        occurrence.activity_kind = "project"
        if project["batch"]:
            occurrence.batch_codes = [project["batch"].upper()]
        if not mappings["project_activity_approved"]:
            occurrence.issues.append("PROJECT activity requires approval")
    elif simple:
        resolve_subject(simple["subject"], mappings, occurrence)
        resolve_faculty(simple["faculty"], known_initials, occurrence)
        occurrence.classroom = simple["room"].upper()
        occurrence.class_ids = [default_class]
    else:
        occurrence.issues.append("unrecognized timetable cell format")
        occurrence.class_ids = [default_class]
    return occurrence


def parse_timetable(workbook_path: Path, mapping_path: Path | None = None) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=False)
    worksheet = workbook["Timetable"]
    mappings = load_mappings(mapping_path)
    known_initials = faculty_initials(workbook)
    occurrences: list[ParsedOccurrence] = []
    current_day: str | None = None

    for row_number in range(1, worksheet.max_row + 1):
        day_match = DAY_PATTERN.match(clean(worksheet.cell(row_number, 5).value))
        if day_match:
            current_day = day_match.group(1).title()
            continue
        row_time_range = parse_time(worksheet.cell(row_number, 5).value)
        if not current_day or not row_time_range:
            continue
        for column, default_class in CLASS_COLUMNS.items():
            cell = worksheet.cell(row_number, column)
            raw = clean(cell.value)
            if not raw:
                continue
            time_range = cell_time_range(worksheet, row_number, column) or row_time_range
            items = split_items(raw)
            for item_index, item in enumerate(items):
                occurrences.append(
                    parse_item(
                        item,
                        day=current_day,
                        time_range=time_range,
                        coordinate=cell.coordinate,
                        item_index=item_index,
                        default_class=default_class,
                        item_count=len(items),
                        mappings=mappings,
                        known_initials=known_initials,
                    )
                )

    source_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    exceptions = [
        {
            "sourceCell": occurrence.source_cell,
            "sourceItemIndex": occurrence.source_item_index,
            "raw": occurrence.raw_source_text,
            "issues": occurrence.issues,
        }
        for occurrence in occurrences
        if occurrence.issues
    ]
    return {
        "sourceWorkbookSha256": source_hash,
        "occurrenceCount": len(occurrences),
        "validOccurrenceCount": sum(not occurrence.issues for occurrence in occurrences),
        "exceptionCount": len(exceptions),
        "strictImportReady": not exceptions,
        "exceptions": exceptions,
        "occurrences": [asdict(occurrence) for occurrence in occurrences],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", required=True, type=Path)
    parser.add_argument("--mappings", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    report = parse_timetable(args.file.resolve(), args.mappings.resolve() if args.mappings else None)
    if args.output:
        args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    printable = {key: value for key, value in report.items() if key != "occurrences"}
    print(json.dumps(printable, indent=2, ensure_ascii=False))
    return 0 if report["strictImportReady"] else 2


if __name__ == "__main__":
    raise SystemExit(main())

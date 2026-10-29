"""Transactional, idempotent SmartClass identity import from data.xlsx.

Timetable rows are deliberately not applied until the review tokens reported by
``validate_college_excel.py`` have approved mappings. The script never deletes
records and never resets an existing password on a repeat run.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import secrets
import string
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
BACKEND_DIR = REPOSITORY_ROOT / "backend"
sys.path.insert(0, str(BACKEND_DIR))

import models  # noqa: E402
from auth import generate_salt, hash_password  # noqa: E402
from database import SessionLocal  # noqa: E402
from validate_college_excel import (  # noqa: E402
    STUDENT_SHEETS,
    clean,
    validate_workbook,
)
from timetable_parser import parse_timetable  # noqa: E402


DEPARTMENT = "Computer Engineering"
COURSES = {
    "SSD": "Software System Design",
    "CCL": "Cloud Computing Lab",
    "DE": "Data Engineering",
    "ML": "Machine Learning",
    "CSS": "Cryptography and System Security",
    "DWM": "Data Warehousing and Mining",
    "SEA": "Software Engineering and Architecture",
    "TCS": "Theory of Computer Science",
}


def excel_text(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean(value)


def normalized_roll(value: Any) -> str:
    roll = excel_text(value)
    if not roll or not re.fullmatch(r"[A-Za-z0-9_-]+", roll):
        raise ValueError("Roll numbers must be non-empty and contain only letters, numbers, '_' or '-'")
    return roll


def student_id(class_id: str, roll_number: str) -> str:
    normalized = roll_number.upper()
    if normalized.isdigit():
        normalized = normalized.zfill(3)
    return f"STU-{class_id}-{normalized}"


def faculty_id(initial: str) -> str:
    normalized = re.sub(r"[^A-Z0-9]", "", initial.upper())
    if not normalized:
        raise ValueError("Faculty initial cannot be normalized into an account ID")
    return f"FAC-{normalized}"


def temporary_password(length: int = 20) -> str:
    alphabet = string.ascii_letters + string.digits + "!@#$%&*+-_"
    required = [
        secrets.choice(string.ascii_uppercase),
        secrets.choice(string.ascii_lowercase),
        secrets.choice(string.digits),
        secrets.choice("!@#$%&*+-_"),
    ]
    required.extend(secrets.choice(alphabet) for _ in range(length - len(required)))
    secrets.SystemRandom().shuffle(required)
    return "".join(required)


def ensure_outside_repository(path: Path) -> Path:
    resolved = path.expanduser().resolve()
    try:
        resolved.relative_to(REPOSITORY_ROOT)
    except ValueError:
        return resolved
    raise ValueError("Credential output must be outside the repository")


def write_credentials(path: Path, credentials: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", newline="", encoding="utf-8") as credential_file:
        try:
            os.chmod(path, 0o600)
        except OSError:
            pass
        writer = csv.DictWriter(
            credential_file,
            fieldnames=["user_id", "temporary_password", "role"],
        )
        writer.writeheader()
        writer.writerows(credentials)


def parse_students(workbook: Any) -> list[dict[str, Any]]:
    students: list[dict[str, Any]] = []
    for sheet_name in STUDENT_SHEETS:
        worksheet = workbook[sheet_name]
        for row_number in range(2, worksheet.max_row + 1):
            values = [worksheet.cell(row_number, column).value for column in range(1, 6)]
            if not any(clean(value) for value in values):
                continue
            email, name, class_value, phone, roll_value = values
            roll_number = normalized_roll(roll_value)
            class_id = clean(class_value)
            if class_id != sheet_name:
                raise ValueError(f"{sheet_name}!C{row_number} does not match its worksheet class")
            students.append(
                {
                    "id": student_id(class_id, roll_number),
                    "name": clean(name),
                    "email": clean(email).casefold(),
                    "phone": excel_text(phone) or None,
                    "roll_no": roll_number,
                    "class_id": class_id,
                    "source_record_key": f"college-xlsx:student:{class_id}:{roll_number.casefold()}",
                }
            )
    return students


def parse_faculty(workbook: Any) -> list[dict[str, Any]]:
    worksheet = workbook["Faculty"]
    headers = [clean(worksheet.cell(1, column).value).casefold() for column in range(1, worksheet.max_column + 1)]
    name_column = headers.index("faculty") + 1
    initial_column = headers.index("initial") + 1
    faculty: list[dict[str, Any]] = []
    for row_number in range(2, worksheet.max_row + 1):
        name = clean(worksheet.cell(row_number, name_column).value)
        initial = clean(worksheet.cell(row_number, initial_column).value).upper()
        if not name and not initial:
            continue
        faculty.append(
            {
                "id": faculty_id(initial),
                "name": name,
                "initial": initial,
                "is_hod": "nupurgiri" in re.sub(r"[^a-z]", "", name.casefold()),
                "source_record_key": f"college-xlsx:faculty:{initial.casefold()}",
            }
        )
    return faculty


def upsert_user(db: Any, record: dict[str, Any], role: str) -> tuple[models.User, str | None]:
    user = db.query(models.User).filter(models.User.id == record["id"]).first()
    password: str | None = None
    if user is None:
        password = temporary_password()
        salt = generate_salt()
        user = models.User(
            id=record["id"],
            password_hash=hash_password(password, salt),
            salt=salt,
            first_login=True,
            token_version=0,
            failed_login_attempts=0,
            source="college-xlsx",
            source_record_key=record["source_record_key"],
        )
        db.add(user)
    elif user.source_record_key not in (None, record["source_record_key"]):
        raise ValueError(f"Account ID collision for {record['id']}")

    user.name = record["name"]
    user.role = role
    user.department = DEPARTMENT
    user.source = "college-xlsx"
    user.source_record_key = record["source_record_key"]
    if role == "student":
        user.email = record["email"]
        user.phone = record["phone"]
        user.class_id = record["class_id"]
        user.is_hod = False
    else:
        # The workbook has no faculty email/phone fields; never invent them.
        user.is_hod = record["is_hod"]
    return user, password


def apply_import(
    workbook_path: Path,
    credential_output: Path,
    validation: dict[str, Any],
    timetable_report: dict[str, Any] | None = None,
) -> dict[str, int]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=False)
    students = parse_students(workbook)
    faculty = parse_faculty(workbook)
    credentials: list[dict[str, str]] = []
    counts = {
        "classes_upserted": 0,
        "courses_upserted": 0,
        "students_upserted": 0,
        "faculty_upserted": 0,
        "accounts_created": 0,
        "schedule_occurrences_upserted": 0,
        "class_course_assignments_upserted": 0,
        "classrooms_upserted": 0,
    }
    credential_written = False
    db = SessionLocal()
    try:
        with db.begin():
            for class_id in STUDENT_SHEETS:
                class_record = db.query(models.ClassModel).filter(models.ClassModel.id == class_id).first()
                if not class_record:
                    class_record = models.ClassModel(id=class_id)
                    db.add(class_record)
                class_record.name = f"Computer Engineering - {class_id}"
                class_record.department = DEPARTMENT
                class_record.is_active = True
                counts["classes_upserted"] += 1

            for code, name in COURSES.items():
                course = db.query(models.CourseModel).filter(models.CourseModel.id == code).first()
                if not course:
                    course = models.CourseModel(id=code)
                    db.add(course)
                course.code = code
                course.name = name
                course.department = DEPARTMENT
                course.is_active = True
                counts["courses_upserted"] += 1

            for record in students:
                user, password = upsert_user(db, record, "student")
                db.flush()
                student = db.query(models.StudentRecord).filter(models.StudentRecord.id == user.id).first()
                if not student:
                    student = models.StudentRecord(id=user.id)
                    db.add(student)
                student.name = record["name"]
                student.roll_no = record["roll_no"]
                student.class_id = record["class_id"]
                student.source = "college-xlsx"
                student.source_record_key = record["source_record_key"]
                counts["students_upserted"] += 1
                if password:
                    counts["accounts_created"] += 1
                    credentials.append({"user_id": user.id, "temporary_password": password, "role": "student"})

            for record in faculty:
                user, password = upsert_user(db, record, "faculty")
                counts["faculty_upserted"] += 1
                if password:
                    counts["accounts_created"] += 1
                    credentials.append({"user_id": user.id, "temporary_password": password, "role": "faculty"})
            # Save pending faculty and student records before timetable references them
            db.flush()
            if timetable_report:
                workbook_hash = timetable_report["sourceWorkbookSha256"]
                for parsed in timetable_report["occurrences"]:
                    occurrence = db.query(models.ScheduleOccurrence).filter(
                        models.ScheduleOccurrence.source_workbook_sha256 == workbook_hash,
                        models.ScheduleOccurrence.source_cell == parsed["source_cell"],
                        models.ScheduleOccurrence.source_item_index == parsed["source_item_index"],
                    ).first()
                    if not occurrence:
                        occurrence = models.ScheduleOccurrence(
                            source_workbook_sha256=workbook_hash,
                            source_sheet="Timetable",
                            source_cell=parsed["source_cell"],
                            source_item_index=parsed["source_item_index"],
                        )
                        db.add(occurrence)
                    occurrence.day_of_week = parsed["day"]
                    occurrence.start_time = parsed["start_time"]
                    occurrence.end_time = parsed["end_time"]
                    occurrence.course_id = parsed["course_id"]
                    occurrence.source_subject_code = parsed["source_subject_code"]
                    occurrence.faculty_id = parsed["faculty_id"]
                    occurrence.classroom = parsed["classroom"]
                    occurrence.activity_kind = parsed["activity_kind"]
                    occurrence.session_type = "Regular"
                    occurrence.parallel_group = parsed["parallel_group"]
                    occurrence.raw_source_text = parsed["raw_source_text"]
                    # Flush only after all non-null occurrence fields are set.
                    # This is required for new rows before membership records
                    # can reference the generated primary key.
                    db.flush()
                    if parsed["classroom"]:
                        classroom = db.query(models.Classroom).filter(
                            models.Classroom.id == parsed["classroom"]
                        ).first()
                        if not classroom:
                            db.add(
                                models.Classroom(
                                    id=parsed["classroom"],
                                    name=parsed["classroom"],
                                    building=None,
                                    capacity=None,
                                    facilities_json="[]",
                                )
                            )
                            counts["classrooms_upserted"] += 1

                    existing_classes = {
                        row.class_id for row in db.query(models.ScheduleClassMembership).filter(
                            models.ScheduleClassMembership.occurrence_id == occurrence.id
                        ).all()
                    }
                    desired_classes = set(parsed["class_ids"])
                    if existing_classes and existing_classes != desired_classes:
                        raise ValueError(
                            f"Approved class mapping changed for {parsed['source_cell']} item "
                            f"{parsed['source_item_index']}; review before altering imported schedule data"
                        )
                    for class_id in sorted(desired_classes - existing_classes):
                        db.add(models.ScheduleClassMembership(occurrence_id=occurrence.id, class_id=class_id))

                    existing_batches = {
                        (row.class_id, row.batch_code)
                        for row in db.query(models.ScheduleBatchGroup).filter(
                            models.ScheduleBatchGroup.occurrence_id == occurrence.id
                        ).all()
                    }
                    desired_batches = {
                        (class_id, batch)
                        for class_id in desired_classes
                        for batch in parsed["batch_codes"]
                    }
                    if existing_batches and existing_batches != desired_batches:
                        raise ValueError(
                            f"Approved batch mapping changed for {parsed['source_cell']} item "
                            f"{parsed['source_item_index']}; review before altering imported schedule data"
                        )
                    for class_id, batch in sorted(desired_batches - existing_batches):
                        db.add(
                            models.ScheduleBatchGroup(
                                occurrence_id=occurrence.id,
                                class_id=class_id,
                                batch_code=batch,
                            )
                        )
                    counts["schedule_occurrences_upserted"] += 1

                    if parsed["course_id"] and parsed["faculty_id"]:
                        for class_id in desired_classes:
                            assignment = db.query(models.ClassCourse).filter(
                                models.ClassCourse.class_id == class_id,
                                models.ClassCourse.course_id == parsed["course_id"],
                                models.ClassCourse.faculty_id == parsed["faculty_id"],
                            ).first()
                            if not assignment:
                                db.add(
                                    models.ClassCourse(
                                        class_id=class_id,
                                        course_id=parsed["course_id"],
                                        faculty_id=parsed["faculty_id"],
                                    )
                                )
                                counts["class_course_assignments_upserted"] += 1

            db.add(
                models.ImportRun(
                    id=str(uuid.uuid4()),
                    workbook_sha256=validation["source"]["sha256"],
                    source_filename=workbook_path.name,
                    status=(
                        "identity-timetable-imported"
                        if timetable_report
                        else "identity-pending-review"
                    ),
                    summary_json=json.dumps(counts, sort_keys=True),
                    completed_at=datetime.now(timezone.utc).replace(tzinfo=None),
                )
            )
            if credentials:
                write_credentials(credential_output, credentials)
                credential_written = True
    except Exception:
        if credential_written:
            credential_output.unlink(missing_ok=True)
        raise
    finally:
        db.close()
    return counts


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", required=True, type=Path)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--dry-run", action="store_true", help="Validate only (default)")
    mode.add_argument("--apply", action="store_true", help="Apply identity records transactionally")
    parser.add_argument(
        "--credential-output",
        type=Path,
        help="New credential CSV outside the repository; required with --apply",
    )
    parser.add_argument(
        "--timetable-mappings",
        type=Path,
        help="Approved mapping JSON; when supplied, timetable import is strict and all exceptions must resolve",
    )
    args = parser.parse_args()

    workbook_path = args.file.resolve()
    validation = validate_workbook(workbook_path)
    timetable_report = (
        parse_timetable(workbook_path, args.timetable_mappings.resolve())
        if args.timetable_mappings
        else None
    )
    identity_issues = [
        issue for issue in validation["issues"] if not issue.startswith("Timetable:")
    ]
    summary = {
        "students_validated": validation["students"]["summary"]["total_records"],
        "faculty_validated": validation["faculty"]["records"],
        "identity_issues": identity_issues,
        "timetable_review_issues": [
            issue for issue in validation["issues"] if issue.startswith("Timetable:")
        ],
    }
    if timetable_report:
        # The workbook validator flags tokens that require a decision. Once an
        # approved mapping resolves them, report only parser exceptions so a
        # successful strict run does not retain stale warnings.
        summary["timetable_review_issues"] = timetable_report["exceptions"]
        summary["timetable_occurrences"] = timetable_report["occurrenceCount"]
        summary["timetable_exceptions"] = timetable_report["exceptionCount"]
        if not timetable_report["strictImportReady"]:
            print(json.dumps(summary, indent=2))
            return 2
    if identity_issues:
        print(json.dumps(summary, indent=2))
        return 2
    if not args.apply:
        summary["mode"] = "dry-run"
        print(json.dumps(summary, indent=2))
        return 0
    if not args.credential_output:
        parser.error("--credential-output is required with --apply")
    credential_output = ensure_outside_repository(args.credential_output)
    counts = apply_import(workbook_path, credential_output, validation, timetable_report)
    summary.update(counts)
    summary["mode"] = "applied"
    summary["credential_manifest"] = (
        "created outside repository"
        if counts["accounts_created"]
        else "not created because no new accounts were provisioned"
    )
    print(json.dumps(summary, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

"""Validate and idempotently import student PEC selections from pec.xlsx."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DECISIONS_PATH = REPOSITORY_ROOT / "config" / "pec_mappings.approved.json"
sys.path.insert(0, str(REPOSITORY_ROOT / "backend"))

import models  # noqa: E402
from database import SessionLocal  # noqa: E402


PEC_MAPPINGS = {
    "CRYPTOGRAPHY AND SYSTEM SECURITY": ("CSS", "A"),
    "SYSTEM DESIGN": ("SSD", "B"),
    "SOFTWARE SYSTEM DESIGN": ("SSD", "B"),
    "DATA WAREHOUSING AND MINING": ("DWM", "C"),
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value)).strip()


def normalized_name(value: Any) -> str:
    return clean(value).casefold()


def read_source(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        expected_class = sheet_name.upper()
        for row_number, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if row_number == 1 or not clean(values[0] if values else None):
                continue
            padded = tuple(values) + (None,) * max(0, 5 - len(values))
            rows.append({
                "sheet": sheet_name,
                "row": row_number,
                "name": clean(padded[0]),
                "class_id": clean(padded[1]).upper(),
                "roll_no": clean(padded[3]),
                "pec_name": clean(padded[4]).upper(),
                "expected_class": expected_class,
            })
    return rows


def read_decisions(path: Path | None) -> tuple[dict[tuple[str, str], dict[str, str]], str | None]:
    if path is None or not path.exists():
        return {}, None
    payload = json.loads(path.read_text(encoding="utf-8"))
    expected_batches = {course_id: batch for course_id, batch in PEC_MAPPINGS.values()}
    declared_batches = payload.get("pecBatchByCourse", {})
    if declared_batches != expected_batches:
        raise ValueError(
            "Approved PEC batch mapping must exactly match CSS=A, SSD=B, DWM=C"
        )
    decisions: dict[tuple[str, str], dict[str, str]] = {}
    for item in payload.get("studentOverrides", []):
        class_id = clean(item.get("classId")).upper()
        roll_no = clean(item.get("rollNo"))
        course_id = clean(item.get("courseId")).upper()
        batch_code = clean(item.get("batchCode")).upper()
        if not class_id or not roll_no or expected_batches.get(course_id) != batch_code:
            raise ValueError(f"Invalid approved PEC override: {item!r}")
        key = (class_id, roll_no)
        if key in decisions:
            raise ValueError(f"Duplicate approved PEC override for {class_id} roll {roll_no}")
        decisions[key] = {
            "courseId": course_id,
            "batchCode": batch_code,
            "reason": clean(item.get("reason")),
        }
    return decisions, str(path)


def validate_and_import(
    path: Path, *, apply: bool, decisions_path: Path | None = DEFAULT_DECISIONS_PATH
) -> dict[str, Any]:
    decisions, decisions_source = read_decisions(decisions_path)
    source_rows = read_source(path)
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    structural_issues: list[dict[str, Any]] = []
    for row in source_rows:
        if row["class_id"] != row["expected_class"]:
            structural_issues.append({
                "type": "class_sheet_mismatch", "sheet": row["sheet"], "row": row["row"]
            })
            continue
        if not row["roll_no"]:
            structural_issues.append({
                "type": "missing_roll_number", "sheet": row["sheet"], "row": row["row"]
            })
            continue
        if row["pec_name"] not in PEC_MAPPINGS:
            structural_issues.append({
                "type": "unknown_pec", "sheet": row["sheet"], "row": row["row"],
                "value": row["pec_name"],
            })
            continue
        grouped[(row["class_id"], row["roll_no"])].append(row)

    report: dict[str, Any] = {
        "mode": "apply" if apply else "dry-run",
        "source": str(path),
        "decisionsSource": decisions_source,
        "sourceRows": len(source_rows),
        "uniqueStudentKeys": len(grouped),
        "validEnrollments": 0,
        "created": 0,
        "existing": 0,
        "duplicateRows": [],
        "conflicts": [],
        "resolvedConflicts": [],
        "unmatched": [],
        "nameMismatches": [],
        "structuralIssues": structural_issues,
        "timetableBatchWarnings": [],
    }

    db = SessionLocal()
    try:
        for (class_id, roll_no), records in sorted(grouped.items()):
            distinct_courses = {PEC_MAPPINGS[row["pec_name"]][0] for row in records}
            if len(distinct_courses) > 1:
                decision = decisions.get((class_id, roll_no))
                if not decision or decision["courseId"] not in distinct_courses:
                    report["conflicts"].append({
                        "classId": class_id,
                        "rollNo": roll_no,
                        "rows": [f'{row["sheet"]}!{row["row"]}' for row in records],
                        "courses": sorted(distinct_courses),
                    })
                    continue
                records = [
                    row for row in records
                    if PEC_MAPPINGS[row["pec_name"]][0] == decision["courseId"]
                ]
                report["resolvedConflicts"].append({
                    "classId": class_id,
                    "rollNo": roll_no,
                    "selectedCourse": decision["courseId"],
                    "selectedBatch": decision["batchCode"],
                    "reason": decision["reason"],
                })
            if len(records) > 1:
                report["duplicateRows"].append({
                    "classId": class_id,
                    "rollNo": roll_no,
                    "rows": [f'{row["sheet"]}!{row["row"]}' for row in records],
                })
            row = records[0]
            student = db.query(models.StudentRecord).filter(
                models.StudentRecord.class_id == class_id,
                models.StudentRecord.roll_no == roll_no,
            ).first()
            if not student:
                report["unmatched"].append({
                    "classId": class_id, "rollNo": roll_no, "sourceRow": f'{row["sheet"]}!{row["row"]}'
                })
                continue
            if normalized_name(student.name) != normalized_name(row["name"]):
                report["nameMismatches"].append({
                    "studentId": student.id, "sourceRow": f'{row["sheet"]}!{row["row"]}'
                })
                continue
            course_id, batch_code = PEC_MAPPINGS[row["pec_name"]]
            if not db.get(models.CourseModel, course_id):
                report["structuralIssues"].append({
                    "type": "course_missing", "courseId": course_id, "sourceRow": f'{row["sheet"]}!{row["row"]}'
                })
                continue
            existing = db.query(models.ElectiveEnrollment).filter_by(student_id=student.id).first()
            if existing:
                if existing.course_id != course_id or existing.batch_code != batch_code:
                    report["conflicts"].append({
                        "studentId": student.id,
                        "existingCourse": existing.course_id,
                        "sourceCourse": course_id,
                        "sourceRow": f'{row["sheet"]}!{row["row"]}',
                    })
                    continue
                report["existing"] += 1
                report["validEnrollments"] += 1
                continue
            report["validEnrollments"] += 1
            if apply:
                db.add(models.ElectiveEnrollment(
                    student_id=student.id,
                    academic_class_id=class_id,
                    course_id=course_id,
                    batch_code=batch_code,
                    source="pec_workbook",
                    source_sheet=row["sheet"],
                    source_row=row["row"],
                    source_record_key=f"PEC:{class_id}:{roll_no}",
                ))
                report["created"] += 1
        if apply:
            db.commit()
        else:
            db.rollback()
        expected_batch = {course_id: batch for course_id, batch in PEC_MAPPINGS.values()}
        timetable_batches = db.query(models.ScheduleOccurrence, models.ScheduleBatchGroup).join(
            models.ScheduleBatchGroup,
            models.ScheduleBatchGroup.occurrence_id == models.ScheduleOccurrence.id,
        ).filter(models.ScheduleOccurrence.course_id.in_(expected_batch)).all()
        report["timetableBatchWarnings"] = [{
            "occurrenceId": occurrence.id,
            "sourceCell": occurrence.source_cell,
            "classId": batch.class_id,
            "courseId": occurrence.course_id,
            "scheduledBatch": batch.batch_code,
            "expectedPecBatch": expected_batch[occurrence.course_id],
        } for occurrence, batch in timetable_batches if batch.batch_code != expected_batch[occurrence.course_id]]
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()

    report["strictImportReady"] = not any([
        report["structuralIssues"], report["conflicts"], report["unmatched"], report["nameMismatches"]
    ])
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", required=True, type=Path)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--output", type=Path)
    parser.add_argument(
        "--decisions",
        type=Path,
        default=DEFAULT_DECISIONS_PATH,
        help="Approved conflict decisions (defaults to config/pec_mappings.approved.json)",
    )
    args = parser.parse_args()
    report = validate_and_import(
        args.file.resolve(), apply=args.apply, decisions_path=args.decisions.resolve()
    )
    if args.output:
        args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    return 0 if report["strictImportReady"] else 2


if __name__ == "__main__":
    raise SystemExit(main())

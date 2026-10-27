"""Read-only validation for the SmartClass college workbook.

The default output is privacy-safe: it reports counts and source row/cell
coordinates without printing student names, email addresses, or phone numbers.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


STUDENT_SHEETS = ("D12A", "D12B", "D12C")
REQUIRED_SHEETS = (*STUDENT_SHEETS, "Faculty", "Timetable")
EXPECTED_STUDENT_HEADERS = (
    "Email Address",
    "FULL NAME (SURNAME FIRST)",
    "Class",
    "Mobile Number",
    "Roll No",
)
EMAIL_PATTERN = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def canonical(value: Any) -> str:
    return clean(value).casefold()


def workbook_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as workbook_file:
        for chunk in iter(lambda: workbook_file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def nonempty_rows(worksheet: Any, start_row: int, columns: int) -> list[tuple[int, list[Any]]]:
    rows: list[tuple[int, list[Any]]] = []
    for row_number in range(start_row, worksheet.max_row + 1):
        values = [worksheet.cell(row_number, column).value for column in range(1, columns + 1)]
        if any(clean(value) for value in values):
            rows.append((row_number, values))
    return rows


def duplicate_coordinates(values: list[tuple[str, str]]) -> dict[str, list[str]]:
    grouped: dict[str, list[str]] = defaultdict(list)
    for normalized_value, coordinate in values:
        if normalized_value:
            grouped[normalized_value].append(coordinate)
    return {
        value: coordinates
        for value, coordinates in grouped.items()
        if len(coordinates) > 1
    }


def validate_students(workbook: Any) -> tuple[dict[str, Any], list[str]]:
    report: dict[str, Any] = {}
    issues: list[str] = []
    workbook_emails: list[tuple[str, str]] = []
    total_records = 0

    for sheet_name in STUDENT_SHEETS:
        worksheet = workbook[sheet_name]
        headers = tuple(clean(worksheet.cell(1, column).value) for column in range(1, 6))
        rows = nonempty_rows(worksheet, 2, 5)
        total_records += len(rows)
        missing_by_column: dict[str, list[int]] = {
            header: [] for header in EXPECTED_STUDENT_HEADERS
        }
        class_values: Counter[str] = Counter()
        email_values: list[tuple[str, str]] = []
        roll_values: list[tuple[str, str]] = []
        malformed_email_rows: list[int] = []
        class_mismatch_rows: list[int] = []

        if headers != EXPECTED_STUDENT_HEADERS:
            issues.append(f"{sheet_name}: unexpected headers {headers!r}")

        for row_number, values in rows:
            for index, header in enumerate(EXPECTED_STUDENT_HEADERS):
                if not clean(values[index]):
                    missing_by_column[header].append(row_number)

            email = canonical(values[0])
            class_name = clean(values[2])
            roll_number = canonical(values[4])
            coordinate_prefix = f"{sheet_name}!"
            if email:
                email_values.append((email, f"{coordinate_prefix}A{row_number}"))
                workbook_emails.append((email, f"{coordinate_prefix}A{row_number}"))
                if not EMAIL_PATTERN.fullmatch(email):
                    malformed_email_rows.append(row_number)
            if roll_number:
                roll_values.append((roll_number, f"{coordinate_prefix}E{row_number}"))
            if class_name:
                class_values[class_name] += 1
            if class_name != sheet_name:
                class_mismatch_rows.append(row_number)

        duplicate_emails = duplicate_coordinates(email_values)
        duplicate_rolls = duplicate_coordinates(roll_values)
        if duplicate_emails:
            issues.append(f"{sheet_name}: duplicate email values detected")
        if duplicate_rolls:
            issues.append(f"{sheet_name}: duplicate roll numbers detected")
        if class_mismatch_rows:
            issues.append(f"{sheet_name}: class value differs from sheet name")
        if malformed_email_rows:
            issues.append(f"{sheet_name}: malformed email values detected")

        report[sheet_name] = {
            "worksheet_rows": worksheet.max_row,
            "records": len(rows),
            "blank_rows_ignored": worksheet.max_row - 1 - len(rows),
            "headers": list(headers),
            "missing_rows_by_column": missing_by_column,
            "class_value_counts": dict(sorted(class_values.items())),
            "class_mismatch_rows": class_mismatch_rows,
            "malformed_email_rows": malformed_email_rows,
            "duplicate_email_coordinates": list(duplicate_emails.values()),
            "duplicate_roll_coordinates": list(duplicate_rolls.values()),
        }

    cross_sheet_duplicates = duplicate_coordinates(workbook_emails)
    if cross_sheet_duplicates:
        issues.append("Student email addresses are duplicated across class sheets")
    report["summary"] = {
        "total_records": total_records,
        "cross_sheet_duplicate_email_coordinates": list(cross_sheet_duplicates.values()),
    }
    return report, issues


def validate_faculty(workbook: Any) -> tuple[dict[str, Any], list[str]]:
    worksheet = workbook["Faculty"]
    rows = nonempty_rows(worksheet, 2, worksheet.max_column)
    headers = [clean(worksheet.cell(1, column).value) for column in range(1, worksheet.max_column + 1)]
    normalized_headers = [header.casefold() for header in headers]
    try:
        name_index = normalized_headers.index("faculty")
        initial_index = normalized_headers.index("initial")
    except ValueError as exc:
        raise ValueError(
            "Faculty worksheet must contain 'Faculty' and 'Initial' headers"
        ) from exc
    initials: list[tuple[str, str]] = []
    missing_name_rows: list[int] = []
    missing_initial_rows: list[int] = []
    hod_matches: list[int] = []

    for row_number, values in rows:
        name = clean(values[name_index])
        initial = clean(values[initial_index]).upper()
        if not name:
            missing_name_rows.append(row_number)
        if not initial:
            missing_initial_rows.append(row_number)
        else:
            initial_column = worksheet.cell(1, initial_index + 1).column_letter
            initials.append((initial.casefold(), f"Faculty!{initial_column}{row_number}"))
        normalized_name = re.sub(r"[^a-z]", "", name.casefold())
        if "nupurgiri" in normalized_name:
            hod_matches.append(row_number)

    duplicates = duplicate_coordinates(initials)
    issues: list[str] = []
    if missing_name_rows:
        issues.append("Faculty: records with missing names detected")
    if missing_initial_rows:
        issues.append("Faculty: records with missing initials detected")
    if duplicates:
        issues.append("Faculty: duplicate initials require stable suffix mapping")
    if len(hod_matches) != 1:
        issues.append("Faculty: expected exactly one Nupur Giri HOD match")

    return {
        "worksheet_rows": worksheet.max_row,
        "records": len(rows),
        "headers": headers,
        "missing_name_rows": missing_name_rows,
        "missing_initial_rows": missing_initial_rows,
        "duplicate_initial_coordinates": list(duplicates.values()),
        "hod_match_rows": hod_matches,
        "initials": [initial.upper() for initial, _ in initials],
    }, issues


def validate_timetable(workbook: Any) -> tuple[dict[str, Any], list[str]]:
    worksheet = workbook["Timetable"]
    populated_cells: list[dict[str, Any]] = []
    for row in worksheet.iter_rows():
        for cell in row:
            raw_value = clean(cell.value)
            if raw_value:
                populated_cells.append({"coordinate": cell.coordinate, "value": raw_value})

    text = "\n".join(item["value"] for item in populated_cells).upper()
    review_tokens = {
        token: [
            item["coordinate"]
            for item in populated_cells
            if re.search(rf"(?<![A-Z0-9]){re.escape(token)}(?![A-Z0-9])", item["value"].upper())
        ]
        for token in ("SD", "D17B", "D17C", "PROJECT")
    }
    review_tokens = {token: cells for token, cells in review_tokens.items() if cells}
    issues = [
        f"Timetable: {token} requires explicit mapping review at {', '.join(cells)}"
        for token, cells in review_tokens.items()
    ]

    return {
        "worksheet_rows": worksheet.max_row,
        "worksheet_columns": worksheet.max_column,
        "populated_cell_count": len(populated_cells),
        "review_tokens": review_tokens,
        "contains_combined_class_marker": "+" in text,
        "populated_cells": populated_cells,
    }, issues


def validate_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    missing_sheets = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
    if missing_sheets:
        raise ValueError(f"Missing required worksheets: {', '.join(missing_sheets)}")

    students, student_issues = validate_students(workbook)
    faculty, faculty_issues = validate_faculty(workbook)
    timetable, timetable_issues = validate_timetable(workbook)
    issues = [*student_issues, *faculty_issues, *timetable_issues]
    return {
        "source": {
            "filename": path.name,
            "sha256": workbook_sha256(path),
            "worksheets": workbook.sheetnames,
        },
        "students": students,
        "faculty": faculty,
        "timetable": timetable,
        "issues": issues,
        "strict_import_ready": not issues,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", type=Path, required=True, help="Path to data.xlsx")
    parser.add_argument("--output", type=Path, help="Optional JSON report destination")
    parser.add_argument(
        "--include-timetable-cells",
        action="store_true",
        help="Include raw timetable cell text in stdout (never includes student records)",
    )
    args = parser.parse_args()

    report = validate_workbook(args.file.resolve())
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    printable = json.loads(json.dumps(report))
    if not args.include_timetable_cells:
        printable["timetable"].pop("populated_cells", None)
    print(json.dumps(printable, indent=2, ensure_ascii=False))
    return 0 if report["strict_import_ready"] else 2


if __name__ == "__main__":
    raise SystemExit(main())

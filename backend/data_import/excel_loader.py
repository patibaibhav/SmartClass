from dataclasses import dataclass
from typing import List, Dict, Optional, Tuple, Any
import openpyxl

@dataclass
class StudentImportData:
    id: str
    name: str
    full_name_raw: str
    email: str
    mobile: Optional[str]
    roll_no: str
    class_id: str
    role: str = "student"
    department: str = "Computer Engineering"

@dataclass
class FacultyImportData:
    id: str
    name: str
    initials: str
    email: str
    role: str = "faculty"
    is_hod: bool = False
    department: str = "Computer Engineering"

def load_students(
    ws_or_wb_path,
    department_name: str = "Computer Engineering"
) -> Tuple[List[StudentImportData], Dict[str, Any]]:
    """Loads students from sheets D12A, D12B, D12C in data.xlsx."""
    if isinstance(ws_or_wb_path, str):
        wb = openpyxl.load_workbook(ws_or_wb_path, data_only=True)
    elif isinstance(ws_or_wb_path, openpyxl.Workbook):
        wb = ws_or_wb_path
    else:
        raise ValueError("Expected workbook path or Workbook instance")

    students: List[StudentImportData] = []
    class_sheets = ["D12A", "D12B", "D12C"]
    stats: Dict[str, Any] = {
        "counts": {},
        "missing_mobiles": {},
        "roll_gaps": {},
        "total": 0
    }

    for sheet_name in class_sheets:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        sheet_students: List[StudentImportData] = []
        rolls_seen = []
        missing_mobile_rolls = []

        for r in range(2, ws.max_row + 1):
            email_cell = ws.cell(r, 1).value
            name_cell = ws.cell(r, 2).value
            class_cell = ws.cell(r, 3).value
            mobile_cell = ws.cell(r, 4).value
            roll_cell = ws.cell(r, 5).value

            if not email_cell and not name_cell and not roll_cell:
                continue

            raw_name = str(name_cell).strip() if name_cell is not None else ""
            cleaned_name = " ".join(raw_name.split()).title()

            email = str(email_cell).strip().lower() if email_cell is not None else ""

            # Mobile null handling: D12A rolls 71 & 72 have no mobile
            mobile: Optional[str] = None
            if mobile_cell is not None and str(mobile_cell).strip():
                m_str = str(mobile_cell).strip()
                # If numeric float/int, convert to int str
                try:
                    mobile = str(int(float(m_str)))
                except ValueError:
                    mobile = m_str
            else:
                mobile = None

            roll_val = int(roll_cell) if roll_cell is not None else r - 1
            rolls_seen.append(roll_val)
            if not mobile:
                missing_mobile_rolls.append(roll_val)

            class_id = str(class_cell).strip() if class_cell is not None else sheet_name
            user_id = f"{class_id}-{roll_val:02d}"

            student = StudentImportData(
                id=user_id,
                name=cleaned_name,
                full_name_raw=raw_name,
                email=email,
                mobile=mobile,
                roll_no=str(roll_val),
                class_id=class_id,
                role="student",
                department=department_name
            )
            sheet_students.append(student)

        # Sort strictly by roll number
        sheet_students.sort(key=lambda s: int(s.roll_no))
        students.extend(sheet_students)

        # Detect gaps
        rolls_seen.sort()
        if rolls_seen:
            full_range = set(range(min(rolls_seen), max(rolls_seen) + 1))
            gaps = sorted(list(full_range - set(rolls_seen)))
        else:
            gaps = []

        stats["counts"][sheet_name] = len(sheet_students)
        stats["missing_mobiles"][sheet_name] = missing_mobile_rolls
        stats["roll_gaps"][sheet_name] = gaps

    stats["total"] = len(students)
    return students, stats


def load_faculty(
    ws_or_wb_path,
    hod_initials: Optional[List[str]] = None,
    department_name: str = "Computer Engineering"
) -> Tuple[List[FacultyImportData], Dict[str, Any]]:
    """Loads all faculty from Faculty sheet in data.xlsx."""
    if isinstance(ws_or_wb_path, str):
        wb = openpyxl.load_workbook(ws_or_wb_path, data_only=True)
    elif isinstance(ws_or_wb_path, openpyxl.Workbook):
        wb = ws_or_wb_path
    else:
        raise ValueError("Expected workbook path or Workbook instance")

    if hod_initials is None:
        hod_initials = []

    ws = wb["Faculty"]
    faculty_list: List[FacultyImportData] = []
    initials_map: Dict[str, str] = {}

    for r in range(2, ws.max_row + 1):
        name_cell = ws.cell(r, 2).value
        init_cell = ws.cell(r, 3).value

        if not name_cell and not init_cell:
            continue

        raw_name = str(name_cell).strip()
        initial = str(init_cell).strip()

        fac_id = f"FAC-{initial}"
        email = f"fac-{initial.lower()}@smartclass.local"
        is_hod = initial in hod_initials

        fac = FacultyImportData(
            id=fac_id,
            name=raw_name,
            initials=initial,
            email=email,
            role="faculty",
            is_hod=is_hod,
            department=department_name
        )
        faculty_list.append(fac)
        initials_map[initial] = raw_name

    stats = {
        "total": len(faculty_list),
        "initials": list(initials_map.keys())
    }
    return faculty_list, stats

import re
from dataclasses import dataclass, field
from typing import List, Dict, Optional, Tuple, Set, Any
import openpyxl

@dataclass
class ParsedTimetableEntry:
    day_of_week: str
    start_time: str
    end_time: str
    time: str
    class_id: str
    course_code: str
    course_name: str
    faculty_initials: Optional[str]
    classroom: Optional[str]
    batch: Optional[str]
    session_kind: str
    group_id: Optional[str]
    is_ambiguous: bool
    source_ref: str

@dataclass
class ParseReport:
    counts_by_day_class: Dict[str, Dict[str, int]]
    total_entries: int
    courses_discovered: List[str]
    rooms_discovered: List[str]
    faculty_initials_discovered: List[str]
    unused_faculty_initials: List[str]
    unknown_faculty_initials: List[str]
    normalisations_applied: List[str]
    ambiguities: List[str]


def parse_slot_time(time_str: str) -> Tuple[Optional[str], Optional[str]]:
    """Convert slot labels like '8:30-9:30', '10.30-11.30', '1:30-2:30' into 24-hour ('08:30', '09:30')."""
    if not time_str:
        return None, None
    s = str(time_str).replace('.', ':').strip()
    m = re.match(r'(\d+):(\d+)\s*-\s*(\d+):(\d+)', s)
    if not m:
        return None, None
    h1, m1, h2, m2 = map(int, m.groups())
    # Slots in schedule: 8:30..12:30 morning, 1:30..3:30 afternoon
    if h1 < 8:
        h1 += 12
    if h2 < 8 or h2 < h1:
        h2 += 12
    return f"{h1:02d}:{m1:02d}", f"{h2:02d}:{m2:02d}"


def parse_timetable_sheet(
    ws_or_wb,
    known_faculty_initials: Optional[Set[str]] = None,
    course_names_map: Optional[Dict[str, str]] = None
) -> Tuple[List[ParsedTimetableEntry], ParseReport]:
    """Pure function parsing the Timetable sheet of data.xlsx."""
    if isinstance(ws_or_wb, str):
        wb = openpyxl.load_workbook(ws_or_wb)
        ws = wb['Timetable']
    elif isinstance(ws_or_wb, openpyxl.Workbook):
        ws = ws_or_wb['Timetable']
    else:
        ws = ws_or_wb

    if course_names_map is None:
        course_names_map = {}

    days_info = [
        ('Monday', 4, 10),
        ('Tuesday', 13, 19),
        ('Wednesday', 22, 28),
        ('Thursday', 31, 37),
        ('Friday', 40, 46)
    ]
    classes_col = {6: 'D12A', 7: 'D12B', 8: 'D12C'}

    # Build row to slot time mapping
    row_times: Dict[int, Tuple[str, str]] = {}
    for day_name, start_row, end_row in days_info:
        for r in range(start_row, end_row + 1):
            tc = ws.cell(r, 5).value
            if tc:
                st, et = parse_slot_time(str(tc))
                if st and et:
                    row_times[r] = (st, et)

    merged_ranges = list(ws.merged_cells.ranges)

    def get_merge_range(r_idx: int, c_idx: int):
        for rng in merged_ranges:
            if r_idx >= rng.min_row and r_idx <= rng.max_row and c_idx >= rng.min_col and c_idx <= rng.max_col:
                return rng
        return None

    entries: List[ParsedTimetableEntry] = []
    ambiguities: List[str] = []
    normalisations: List[str] = []
    courses_found: Set[str] = set()
    rooms_found: Set[str] = set()
    initials_found: Set[str] = set()

    for day_name, start_row, end_row in days_info:
        for r in range(start_row, end_row + 1):
            if r not in row_times:
                continue
            st_slot, et_slot = row_times[r]
            # 12:30-1:30 is lunch and is empty — import nothing
            if st_slot == '12:30':
                continue

            for c in [6, 7, 8]:
                mr = get_merge_range(r, c)
                if mr:
                    # Attribution to top-left value; skip covered cells that are not top-left
                    if r != mr.min_row or c != mr.min_col:
                        continue
                    st = row_times[mr.min_row][0]
                    et = row_times[mr.max_row][1]
                    coord_ref = f"Timetable!{mr.coord}"
                else:
                    st, et = st_slot, et_slot
                    coord_ref = f"Timetable!{ws.cell(r, c).coordinate}"

                val = ws.cell(r, c).value
                if not val or not str(val).strip():
                    continue

                text = str(val).strip()

                raw_items = [i.strip() for i in re.split(r'\n\s*\n', text) if i.strip()]
                items: List[str] = []
                for it in raw_items:
                    sublines = [l.strip() for l in it.splitlines() if l.strip()]
                    # Whole class theory: COURSE, (FAC), ROOM
                    if len(sublines) == 3 and sublines[1].startswith('(') and sublines[1].endswith(')'):
                        items.append('\n'.join(sublines))
                    else:
                        for sl in sublines:
                            items.append(sl)

                cell_batch_counts: Dict[str, int] = {}

                for it in items:
                    it = it.strip()
                    if not it:
                        continue

                    # 1. Joint / elective block (multi-class)
                    # e.g. SD / (PRL) / 315 (D12A+D12C)
                    m_joint = re.match(r'^([A-Z0-9]+)\s*/\s*\(?([A-Z]+)\)?\s*/\s*([A-Z0-9]+)\s*\(([^)]+)\)$', it)
                    if m_joint:
                        crs, fac, rm, cls_str = m_joint.groups()
                        courses_found.add(crs)
                        initials_found.add(fac)
                        rooms_found.add(rm)

                        if 'D17B' in cls_str or 'D17C' in cls_str:
                            orig = cls_str
                            cls_str = cls_str.replace('D17B', 'D12B').replace('D17C', 'D12C')
                            normalisations.append(f"{coord_ref}: normalised typo {orig} -> {cls_str}")

                        target_classes = [x.strip() for x in cls_str.split('+') if x.strip()]
                        gid = f"ELECTIVE_{day_name[:3].upper()}_{st.replace(':', '')}"

                        for tcls in target_classes:
                            is_amb = False
                            # Known ambiguity: D12C assigned to multiple SD sections in joint slot
                            if tcls == 'D12C' and crs == 'SD':
                                prev_sd = [
                                    e for e in entries
                                    if e.day_of_week == day_name and e.start_time == st and e.class_id == 'D12C' and e.course_code == 'SD'
                                ]
                                if prev_sd:
                                    is_amb = True
                                    prev_sd[0].is_ambiguous = True
                                    ambiguities.append(f"{coord_ref}: D12C assigned to multiple SD sections in joint slot {day_name} {st}-{et}")

                            c_name = course_names_map.get(crs, crs)
                            if c_name == "TODO: confirm":
                                c_name = crs

                            entries.append(ParsedTimetableEntry(
                                day_of_week=day_name,
                                start_time=st,
                                end_time=et,
                                time=f"{st}–{et}",
                                class_id=tcls,
                                course_code=crs,
                                course_name=c_name,
                                faculty_initials=fac,
                                classroom=rm,
                                batch=None,
                                session_kind="Theory",
                                group_id=gid,
                                is_ambiguous=is_amb,
                                source_ref=coord_ref
                            ))
                        continue

                    # 2. Project
                    # e.g. PROJECT, A/ PROJECT, B/PROJECT
                    m_proj = re.match(r'^(?:([A-C])\s*/\s*)?PROJECT$', it)
                    if m_proj:
                        pbatch = m_proj.group(1)
                        cur_cls = classes_col[c]
                        courses_found.add("PROJECT")
                        is_amb = False

                        # Known ambiguity: Thursday 1:30 D12C has A/ PROJECT next to batches B and C
                        if day_name == 'Thursday' and st == '13:30' and cur_cls == 'D12C' and pbatch == 'A':
                            is_amb = True
                            ambiguities.append(f"{coord_ref}: Thursday 13:30 D12C has A/ PROJECT alongside batches B and C")

                        if pbatch:
                            cell_batch_counts[pbatch] = cell_batch_counts.get(pbatch, 0) + 1

                        entries.append(ParsedTimetableEntry(
                            day_of_week=day_name,
                            start_time=st,
                            end_time=et,
                            time=f"{st}–{et}",
                            class_id=cur_cls,
                            course_code="PROJECT",
                            course_name="Project",
                            faculty_initials=None,
                            classroom=None,
                            batch=pbatch,
                            session_kind="Project",
                            group_id=None,
                            is_ambiguous=is_amb,
                            source_ref=coord_ref
                        ))
                        continue

                    # 3. Batch-wise practical
                    # e.g. SEA/ A/ 302/ MM
                    m_pract = re.match(r'^([A-Z0-9]+)\s*/\s*([A-C])\s*/\s*([A-Z0-9]+)\s*/\s*\(?([A-Z]+)\)?$', it)
                    if m_pract:
                        crs, bat, rm, fac = m_pract.groups()
                        courses_found.add(crs)
                        initials_found.add(fac)
                        rooms_found.add(rm)

                        cur_cls = classes_col[c]
                        cell_batch_counts[bat] = cell_batch_counts.get(bat, 0) + 1
                        is_amb = False

                        # Known ambiguity: Duplicate batch in same practical slot
                        if cell_batch_counts[bat] > 1:
                            is_amb = True
                            for prev_e in entries:
                                if prev_e.source_ref == coord_ref and prev_e.batch == bat:
                                    prev_e.is_ambiguous = True
                            ambiguities.append(f"{coord_ref}: Duplicate batch {bat} in practical slot {day_name} {st}-{et} ({crs})")

                        c_name = course_names_map.get(crs, crs)
                        if c_name == "TODO: confirm":
                            c_name = crs

                        entries.append(ParsedTimetableEntry(
                            day_of_week=day_name,
                            start_time=st,
                            end_time=et,
                            time=f"{st}–{et}",
                            class_id=cur_cls,
                            course_code=crs,
                            course_name=c_name,
                            faculty_initials=fac,
                            classroom=rm,
                            batch=bat,
                            session_kind="Practical",
                            group_id=None,
                            is_ambiguous=is_amb,
                            source_ref=coord_ref
                        ))
                        continue

                    # 4. Theory, whole class
                    # e.g. COURSE\n(FAC)\nROOM
                    lines = [l.strip() for l in it.splitlines() if l.strip()]
                    if len(lines) == 3:
                        crs = lines[0]
                        fac = re.sub(r'[\(\)]', '', lines[1]).strip()
                        rm = lines[2]
                        courses_found.add(crs)
                        initials_found.add(fac)
                        rooms_found.add(rm)

                        cur_cls = classes_col[c]
                        c_name = course_names_map.get(crs, crs)
                        if c_name == "TODO: confirm":
                            c_name = crs

                        entries.append(ParsedTimetableEntry(
                            day_of_week=day_name,
                            start_time=st,
                            end_time=et,
                            time=f"{st}–{et}",
                            class_id=cur_cls,
                            course_code=crs,
                            course_name=c_name,
                            faculty_initials=fac,
                            classroom=rm,
                            batch=None,
                            session_kind="Theory",
                            group_id=None,
                            is_ambiguous=False,
                            source_ref=coord_ref
                        ))
                        continue

                    # Fallback unparsed
                    ambiguities.append(f"{coord_ref}: Unrecognized timetable format '{it}'")

    # Counts by day and class
    counts_by_day_class: Dict[str, Dict[str, int]] = {}
    for e in entries:
        counts_by_day_class.setdefault(e.day_of_week, {}).setdefault(e.class_id, 0)
        counts_by_day_class[e.day_of_week][e.class_id] += 1

    unknown_initials = []
    unused_initials = []
    if known_faculty_initials:
        unknown_initials = sorted(list(initials_found - known_faculty_initials))
        unused_initials = sorted(list(known_faculty_initials - initials_found))

    report = ParseReport(
        counts_by_day_class=counts_by_day_class,
        total_entries=len(entries),
        courses_discovered=sorted(list(courses_found)),
        rooms_discovered=sorted(list(rooms_found)),
        faculty_initials_discovered=sorted(list(initials_found)),
        unused_faculty_initials=unused_initials,
        unknown_faculty_initials=unknown_initials,
        normalisations_applied=normalisations,
        ambiguities=ambiguities
    )

    return entries, report

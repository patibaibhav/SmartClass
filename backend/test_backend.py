import os
from datetime import datetime
import json
import sys
from pathlib import Path

os.environ.setdefault("APP_ENV", "test")
os.environ.setdefault("DATABASE_URL", "sqlite://")
os.environ.setdefault("SECRET_KEY", "test-only-secret-key-with-at-least-32-characters")

import pytest
from openpyxl import Workbook
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import models
from auth import generate_salt, hash_password
from database import Base, get_db
from main import app

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from timetable_parser import cell_time_range, parse_time  # noqa: E402


test_engine = create_engine(
    "sqlite://",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSession = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


def override_get_db():
    db = TestingSession()
    try:
        yield db
    finally:
        db.close()


app.dependency_overrides[get_db] = override_get_db
client = TestClient(app)


def add_user(
    db,
    user_id: str,
    role: str,
    *,
    class_id: str | None = None,
    is_hod: bool = False,
    first_login: bool = False,
    password: str = "SecurePass!123",
    name: str | None = None,
    source: str = "manual",
):
    salt = generate_salt()
    db.add(
        models.User(
            id=user_id,
            name=name or user_id,
            email=f"{user_id.casefold()}@example.test",
            password_hash=hash_password(password, salt),
            salt=salt,
            role=role,
            is_hod=is_hod,
            department="Computer Engineering",
            class_id=class_id,
            first_login=first_login,
            token_version=0,
            source=source,
        )
    )


@pytest.fixture(autouse=True)
def seeded_database():
    Base.metadata.create_all(test_engine)
    db = TestingSession()
    for class_id in ("D12A", "D12B"):
        db.add(
            models.ClassModel(
                id=class_id,
                name=f"Computer Engineering - {class_id}",
                department="Computer Engineering",
            )
        )
    db.add(
        models.CourseModel(
            id="TCS",
            code="TCS",
            name="Theory of Computer Science",
            department="Computer Engineering",
        )
    )
    add_user(db, "STU-D12A-001", "student", class_id="D12A")
    add_user(db, "STU-D12A-002", "student", class_id="D12A")
    add_user(db, "STU-D12B-001", "student", class_id="D12B")
    add_user(db, "FAC-SM", "faculty")
    add_user(db, "FAC-OTHER", "faculty")
    add_user(db, "FAC-NG", "faculty", is_hod=True)
    add_user(db, "ADM-001", "admin")
    add_user(db, "CON-001", "council")
    add_user(db, "STU-TEMP", "student", class_id="D12A", first_login=True, password="Temporary!123")
    db.flush()
    for student_id, roll, class_id in (
        ("STU-D12A-001", "1", "D12A"),
        ("STU-D12A-002", "2", "D12A"),
        ("STU-D12B-001", "1", "D12B"),
        ("STU-TEMP", "99", "D12A"),
    ):
        db.add(models.StudentRecord(id=student_id, name=student_id, roll_no=roll, class_id=class_id))
    db.add(models.ClassCourse(class_id="D12A", course_id="TCS", faculty_id="FAC-SM"))
    db.add(models.ClassCourse(class_id="D12B", course_id="TCS", faculty_id="FAC-OTHER"))
    db.add(models.Classroom(id="314", name="314", building="Main", capacity=80))
    db.add(
        models.TimetableEntry(
            day_of_week="Monday",
            time="08:30–09:30",
            start_time="08:30",
            end_time="09:30",
            class_id="D12A",
            course_id="TCS",
            course_name="Theory of Computer Science",
            course_code="TCS",
            faculty_id="FAC-SM",
            faculty_name="FAC-SM",
            classroom="314",
        )
    )
    db.commit()
    db.close()
    yield
    Base.metadata.drop_all(test_engine)


def login(user_id: str, password: str = "SecurePass!123") -> str:
    response = client.post(
        "/api/auth/login",
        json={"user_id": user_id, "password": password},
    )
    assert response.status_code == 200, response.text
    return response.json()["token"]


def auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def test_root_and_protected_endpoint():
    assert client.get("/").status_code == 200
    assert client.get("/api/classrooms").status_code == 401


def test_login_uses_generic_error_and_returns_scoped_token():
    response = client.post(
        "/api/auth/login",
        json={"user_id": "STU-D12A-001", "password": "SecurePass!123"},
    )
    assert response.status_code == 200
    assert response.json()["user"]["class"] == "D12A"

    wrong_password = client.post(
        "/api/auth/login",
        json={"user_id": "STU-D12A-001", "password": "wrong"},
    )
    unknown_user = client.post(
        "/api/auth/login",
        json={"user_id": "UNKNOWN", "password": "wrong"},
    )
    assert wrong_password.status_code == unknown_user.status_code == 401
    assert wrong_password.json()["detail"] == unknown_user.json()["detail"]


def test_first_login_requires_limited_token_and_strong_password():
    login_response = client.post(
        "/api/auth/login",
        json={"user_id": "STU-TEMP", "password": "Temporary!123"},
    )
    assert login_response.status_code == 200
    assert login_response.json()["phase"] == "first-login"
    setup_token = login_response.json()["token"]
    assert client.get("/api/classrooms", headers=auth(setup_token)).status_code == 401
    assert client.post(
        "/api/auth/first-login",
        json={"new_password": "weak", "confirm_password": "weak"},
        headers=auth(setup_token),
    ).status_code == 400

    changed = client.post(
        "/api/auth/first-login",
        json={"new_password": "NewSecurePass!456", "confirm_password": "NewSecurePass!456"},
        headers=auth(setup_token),
    )
    assert changed.status_code == 200
    assert client.get("/api/classrooms", headers=auth(changed.json()["token"])).status_code == 200


def test_logout_revokes_token():
    token = login("STU-D12A-001")
    assert client.post("/api/auth/logout", headers=auth(token)).status_code == 200
    assert client.get("/api/auth/me", headers=auth(token)).status_code == 401


def test_student_cannot_view_another_class_timetable():
    token = login("STU-D12A-001")
    own = client.get("/api/classes/D12A/timetable", headers=auth(token))
    other = client.get("/api/classes/D12B/timetable", headers=auth(token))
    assert own.status_code == 200
    assert other.status_code == 403


def test_same_class_students_receive_same_timetable():
    first = client.get("/api/timetable", headers=auth(login("STU-D12A-001")))
    second = client.get("/api/timetable", headers=auth(login("STU-D12A-002")))
    assert first.status_code == second.status_code == 200
    assert first.json() == second.json()
    assert len(first.json()["Monday"]) == 1


def test_conflict_check_is_role_protected_and_detects_overlap():
    query = "day=Monday&start_time=09:00&end_time=10:00&classroom=314&faculty_id=FAC-SM&class_id=D12A"
    student_response = client.post(
        f"/api/timetable/check-conflict?{query}",
        headers=auth(login("STU-D12A-001")),
    )
    admin_response = client.post(
        f"/api/timetable/check-conflict?{query}",
        headers=auth(login("ADM-001")),
    )
    assert student_response.status_code == 403
    assert admin_response.status_code == 200
    assert admin_response.json()["hasConflict"] is True


def test_student_grievances_are_isolated():
    first_token = login("STU-D12A-001")
    second_token = login("STU-D12A-002")
    created = client.post(
        "/api/grievances",
        json={"subject": "Test", "category": "Academic", "description": "Private"},
        headers=auth(first_token),
    )
    assert created.status_code == 200
    assert len(client.get("/api/grievances", headers=auth(first_token)).json()) == 1
    assert client.get("/api/grievances", headers=auth(second_token)).json() == []


def test_attendance_is_calculated_from_conducted_lectures():
    db = TestingSession()
    lecture = models.Lecture(
        class_id="D12A",
        course_id="TCS",
        scheduled_faculty_id="FAC-SM",
        actual_faculty_id="FAC-SM",
        date="2026-10-20",
        start_time="08:30",
        end_time="09:30",
        lecture_type="Regular",
        classroom="314",
        status="Conducted",
    )
    db.add(lecture)
    db.flush()
    db.add(
        models.Attendance(
            lecture_id=lecture.id,
            student_id="STU-D12A-001",
            status="Present",
            marked_by="FAC-SM",
        )
    )
    db.commit()
    db.close()

    response = client.get(
        "/api/attendance/student",
        headers=auth(login("STU-D12A-001")),
    )
    assert response.status_code == 200
    assert response.json()["overallPercentage"] == 100.0
    assert response.json()["courses"][0]["total"] == 1


def test_rule_based_topic_insights_use_multiple_recorded_attempts():
    db = TestingSession()
    for quiz_index, correct_answers in ((1, 2), (2, 3)):
        quiz_id = f"QUIZ-{quiz_index}"
        db.add(
            models.Quiz(
                id=quiz_id,
                course_id="TCS",
                course_name="Theory of Computer Science",
                course_code="TCS",
                class_id="D12A",
                title=quiz_id,
                duration_minutes=15,
                opens_at="2026-10-01 09:00",
                closes_at="2026-10-01 10:00",
                is_released=True,
                created_by="FAC-SM",
            )
        )
        db.flush()
        questions = []
        for question_index in range(3):
            question = models.QuizQuestion(
                quiz_id=quiz_id,
                question_text=f"Question {question_index}",
                options_json='["A", "B"]',
                correct_option=0,
                marks=1,
                topic="Automata",
            )
            db.add(question)
            questions.append(question)
        db.flush()
        answers = {
            str(question.id): 0 if index < correct_answers else 1
            for index, question in enumerate(questions)
        }
        db.add(
            models.QuizAttempt(
                quiz_id=quiz_id,
                student_id="STU-D12A-001",
                started_at=datetime(2026, 10, quiz_index, 9, 0),
                submitted_at=datetime(2026, 10, quiz_index, 9, 10),
                answers_json=json.dumps(answers),
                score=correct_answers,
                total_marks=3,
            )
        )
    db.commit()
    db.close()

    response = client.get(
        "/api/performance-insights?course_id=TCS",
        headers=auth(login("STU-D12A-001")),
    )
    assert response.status_code == 200
    data = response.json()
    assert data["title"] == "Topic-wise Performance Insights (rule-based)"
    assert data["attemptCount"] == 2
    assert data["topics"][0]["accuracy"] == 83.3
    assert data["topics"][0]["classification"] == "strong"
    assert data["trend"]["direction"] == "improving"


def test_forum_replies_require_course_membership():
    db = TestingSession()
    db.add(models.CourseModel(id="SSD", code="SSD", name="Software System Design", department="Computer Engineering"))
    db.add(models.ClassCourse(class_id="D12A", course_id="SSD", faculty_id="FAC-SM"))
    db.add(models.ElectiveEnrollment(
        student_id="STU-D12A-001", academic_class_id="D12A", course_id="SSD", batch_code="B",
        source="test", source_sheet="test", source_row=1, source_record_key="TEST:FORUM:PEC",
    ))
    db.add(models.ForumTopic(
        id="TOPIC-SSD", course_id="SSD", course_code="SSD", title="Scoped",
        content="Only members", author_id="FAC-SM", author_name="FAC-SM",
        created_at="2026-10-27 10:00", replies_count=0, last_reply="Never",
    ))
    db.commit()
    db.close()

    allowed = client.get("/api/forums/TOPIC-SSD/replies", headers=auth(login("STU-D12A-001")))
    denied = client.get("/api/forums/TOPIC-SSD/replies", headers=auth(login("STU-D12B-001")))
    assert allowed.status_code == 200
    assert denied.status_code == 403


def test_quiz_questions_are_hidden_outside_availability_window():
    db = TestingSession()
    db.add(models.Quiz(
        id="FUTURE", course_id="TCS", course_name="Theory of Computer Science",
        course_code="TCS", class_id="D12A", title="Future quiz", duration_minutes=10,
        opens_at="2099-01-01 09:00", closes_at="2099-01-01 10:00",
        is_released=True, created_by="FAC-SM",
    ))
    db.commit()
    db.close()
    response = client.get("/api/quizzes/FUTURE/questions", headers=auth(login("STU-D12A-001")))
    assert response.status_code == 409


def test_faculty_my_classes_are_assignment_scoped():
    response = client.get("/api/academic/my-class", headers=auth(login("FAC-SM")))
    assert response.status_code == 200
    assert [item["classId"] for item in response.json()["classes"]] == ["D12A"]


def test_unassigned_hod_faculty_is_not_given_department_assignments():
    headers = auth(login("FAC-NG"))
    classes = client.get("/api/academic/classes", headers=headers)
    courses = client.get("/api/academic/courses", headers=headers)
    my_classes = client.get("/api/academic/my-class", headers=headers)
    dashboard = client.get("/api/dashboard/stats?role=faculty", headers=headers)
    assert classes.status_code == courses.status_code == 200
    assert my_classes.status_code == dashboard.status_code == 200
    assert classes.json() == []
    assert courses.json() == []
    assert my_classes.json() == {"classes": []}
    assert dashboard.json()["classes"] == 0
    assert dashboard.json()["assignedClasses"] == []
    assert dashboard.json()["assignedCourses"] == []
    assert dashboard.json()["todaySchedule"] == []


def test_time_normalization_and_merged_two_hour_duration():
    assert parse_time("1:30-2:30") == ("13:30", "14:30")
    assert parse_time("12:30-1:30") == ("12:30", "13:30")
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(1, 5).value = "10:30-11:30"
    sheet.cell(2, 5).value = "11:30-12:30"
    sheet.cell(1, 6).value = "TCS (SM) 314"
    sheet.merge_cells("F1:F2")
    assert cell_time_range(sheet, 1, 6) == ("10:30", "12:30")


def test_pec_students_receive_only_their_selected_elective_timetable():
    db = TestingSession()
    for course_id, name in (("CSS", "Cryptography and System Security"), ("SSD", "Software System Design")):
        db.add(models.CourseModel(id=course_id, code=course_id, name=name, department="Computer Engineering"))
        db.add(models.ClassCourse(class_id="D12A", course_id=course_id, faculty_id="FAC-SM"))
        occurrence = models.ScheduleOccurrence(
            day_of_week="Monday", start_time="09:30", end_time="10:30",
            course_id=course_id, source_subject_code=course_id, faculty_id="FAC-SM",
            classroom="314", activity_kind="lecture", session_type="Regular",
            raw_source_text=course_id, source_sheet="test", source_cell=f"{course_id}1",
            source_item_index=0, source_workbook_sha256="a" * 64,
        )
        db.add(occurrence); db.flush()
        db.add(models.ScheduleClassMembership(occurrence_id=occurrence.id, class_id="D12A"))
    practical = models.ScheduleOccurrence(
        day_of_week="Monday", start_time="10:30", end_time="12:30",
        course_id="CSS", source_subject_code="CSS", faculty_id="FAC-SM",
        classroom="314", activity_kind="practical", session_type="Regular",
        raw_source_text="CSS /B/ 314/ SM", source_sheet="test", source_cell="P1",
        source_item_index=0, source_workbook_sha256="b" * 64,
    )
    db.add(practical); db.flush()
    db.add(models.ScheduleClassMembership(occurrence_id=practical.id, class_id="D12A"))
    db.add(models.ScheduleBatchGroup(
        occurrence_id=practical.id, class_id="D12A", batch_code="B"
    ))
    regular_practical = models.ScheduleOccurrence(
        day_of_week="Monday", start_time="12:30", end_time="14:30",
        course_id="TCS", source_subject_code="TCS", faculty_id="FAC-SM",
        classroom="314", activity_kind="practical", session_type="Regular",
        raw_source_text="TCS /B/ 314/ SM", source_sheet="test", source_cell="P2",
        source_item_index=0, source_workbook_sha256="c" * 64,
    )
    db.add(regular_practical); db.flush()
    db.add(models.ScheduleClassMembership(occurrence_id=regular_practical.id, class_id="D12A"))
    db.add(models.ScheduleBatchGroup(
        occurrence_id=regular_practical.id, class_id="D12A", batch_code="B"
    ))
    db.add(models.ElectiveEnrollment(
        student_id="STU-D12A-001", academic_class_id="D12A", course_id="CSS", batch_code="A",
        source="test", source_sheet="test", source_row=1, source_record_key="TEST:PEC:1",
    ))
    db.add(models.ElectiveEnrollment(
        student_id="STU-D12A-002", academic_class_id="D12A", course_id="SSD", batch_code="B",
        source="test", source_sheet="test", source_row=2, source_record_key="TEST:PEC:2",
    ))
    db.commit(); db.close()
    css = client.get("/api/timetable", headers=auth(login("STU-D12A-001"))).json()["Monday"]
    ssd = client.get("/api/timetable", headers=auth(login("STU-D12A-002"))).json()["Monday"]
    assert {row["code"] for row in css} == {"CSS"}
    assert {row["code"] for row in ssd} == {"SSD", "TCS"}
    assert len(css) == 1
    assert len(ssd) == 3
    assert next(row for row in ssd if row["startTime"] == "10:30")["batches"] == ["D12A:B"]


def test_faculty_can_inspect_quiz_questions_and_attempts_without_student_permissions():
    db = TestingSession()
    quiz = models.Quiz(
        id="MANAGED", course_id="TCS", course_name="Theory of Computer Science", course_code="TCS",
        class_id="D12A", title="Managed quiz", duration_minutes=10,
        opens_at="2026-01-01T00:00:00", closes_at="2099-01-01T00:00:00",
        max_attempts=1, is_released=True, created_by="FAC-SM",
    )
    db.add(quiz); db.flush()
    db.add(models.QuizQuestion(quiz_id=quiz.id, question_text="A deterministic question?", options_json='["A","B","C","D"]', correct_option=2, marks=2, topic="Testing"))
    db.commit(); db.close()
    faculty = client.get("/api/quizzes/MANAGED/questions", headers=auth(login("FAC-SM")))
    student = client.get("/api/quizzes/MANAGED/questions", headers=auth(login("STU-D12A-001")))
    assert faculty.status_code == student.status_code == 200
    assert faculty.json()[0]["correctOption"] == 2
    assert "correctOption" not in student.json()[0]
    question_id = student.json()[0]["id"]
    submitted = client.post(
        "/api/quizzes/MANAGED/submit",
        headers=auth(login("STU-D12A-001")),
        json={"answers": {str(question_id): 2}},
    )
    assert submitted.status_code == 200 and submitted.json()["score"] == 2
    attempts = client.get("/api/quizzes/MANAGED/attempts", headers=auth(login("FAC-SM")))
    assert attempts.status_code == 200
    assert attempts.json()[0]["studentId"] == "STU-D12A-001"
    assert attempts.json()[0]["percentage"] == 100.0


def test_cross_role_assignment_submission_and_grading():
    faculty_headers = auth(login("FAC-SM"))
    created = client.post("/api/assignments", headers=faculty_headers, json={
        "course_id": "TCS", "class_id": "D12A", "title": "Integration assignment",
        "description": "Submit a short response", "due_date": "2099-01-01", "marks": 20,
        "allow_late": True,
    })
    assert created.status_code == 200, created.text
    assignment_id = created.json()["assignmentId"]
    submitted = client.post(f"/api/assignments/{assignment_id}/submit", headers=auth(login("STU-D12A-001")), json={"submission_text": "Persisted response"})
    assert submitted.status_code == 200, submitted.text
    submissions = client.get(f"/api/assignments/{assignment_id}/submissions", headers=faculty_headers)
    assert submissions.status_code == 200
    assert submissions.json()[0]["submissionText"] == "Persisted response"
    submission_id = submissions.json()[0]["id"]
    graded = client.post(f"/api/assignments/submissions/{submission_id}/grade", headers=faculty_headers, json={"marks": 18, "feedback": "Good work"})
    assert graded.status_code == 200
    student_items = client.get("/api/assignments", headers=auth(login("STU-D12A-001"))).json()
    item = next(row for row in student_items if row["id"] == assignment_id)
    assert item["status"] == "graded" and item["obtainedMarks"] == 18 and item["feedback"] == "Good work"


def test_cross_role_attendance_edit_updates_student_and_writes_audit():
    faculty_headers = auth(login("FAC-SM"))
    created = client.post("/api/lectures/create", headers=faculty_headers, json={
        "class_id": "D12A", "course_id": "TCS", "actual_faculty_id": "FAC-SM",
        "date": "2026-10-28", "start_time": "09:30", "end_time": "10:30",
        "classroom": "314", "lecture_type": "Regular",
    })
    assert created.status_code == 200, created.text
    lecture_id = created.json()["lectureId"]
    roster = client.get(f"/api/lectures/{lecture_id}/roster", headers=faculty_headers)
    assert roster.status_code == 200 and len(roster.json()["students"]) == 3
    present = client.post("/api/attendance/mark", headers=faculty_headers, json={
        "lecture_id": lecture_id,
        "attendance": [{"student_id": row["id"], "status": "Present"} for row in roster.json()["students"]],
    })
    assert present.status_code == 200
    absent = client.post("/api/attendance/mark", headers=faculty_headers, json={
        "lecture_id": lecture_id,
        "attendance": [{"student_id": row["id"], "status": "Absent" if row["id"] == "STU-D12A-001" else "Present"} for row in roster.json()["students"]],
    })
    assert absent.status_code == 200
    student = client.get("/api/attendance/student", headers=auth(login("STU-D12A-001"))).json()
    assert student["history"][0]["status"] == "Absent"
    db = TestingSession()
    assert db.query(models.AttendanceAudit).filter_by(changed_by="FAC-SM", new_status="Absent").count() == 1
    db.close()


def test_classroom_overlap_allows_adjacent_and_ignores_cancelled():
    db = TestingSession()
    db.add(models.Lecture(
        class_id="D12A", course_id="TCS", scheduled_faculty_id="FAC-SM", actual_faculty_id="FAC-SM",
        date="2026-10-28", start_time="10:00", end_time="12:00", lecture_type="Regular",
        classroom="314", status="Scheduled",
    ))
    db.add(models.Lecture(
        class_id="D12A", course_id="TCS", scheduled_faculty_id="FAC-SM", actual_faculty_id="FAC-SM",
        date="2026-10-28", start_time="12:00", end_time="13:00", lecture_type="Cancelled",
        classroom="314", status="Cancelled",
    ))
    db.commit(); db.close()
    headers = auth(login("STU-D12A-001"))
    adjacent = client.post("/api/classrooms/search-availability", headers=headers, json={"date":"2026-10-28","start_time":"09:00","end_time":"10:00"}).json()
    overlap = client.post("/api/classrooms/search-availability", headers=headers, json={"date":"2026-10-28","start_time":"10:30","end_time":"11:00"}).json()
    cancelled = client.post("/api/classrooms/search-availability", headers=headers, json={"date":"2026-10-28","start_time":"12:00","end_time":"13:00"}).json()
    assert adjacent["classrooms"][0]["status"] == "Available"
    assert overlap["classrooms"][0]["status"] == "Occupied"
    assert cancelled["classrooms"][0]["status"] == "Available"


def test_hod_can_resolve_grievance_and_view_department_dashboard():
    created = client.post(
        "/api/grievances",
        headers=auth(login("STU-D12A-001")),
        json={"subject": "Attendance correction", "category": "Attendance", "description": "Please review"},
    )
    grievance_id = created.json()["grievanceId"]
    hod_headers = auth(login("FAC-NG"))
    listed = client.get("/api/grievances", headers=hod_headers)
    assert listed.status_code == 200
    assert any(item["id"] == grievance_id for item in listed.json())
    updated = client.patch(
        f"/api/grievances/{grievance_id}",
        headers=hod_headers,
        json={"status": "RESOLVED", "resolution": "Attendance record reviewed"},
    )
    assert updated.status_code == 200
    dashboard = client.get("/api/dashboard/stats?role=hod", headers=hod_headers)
    assert dashboard.status_code == 200
    assert dashboard.json()["totalStudents"] == 4
    db = TestingSession()
    assert db.query(models.AuditLog).filter_by(user_id="FAC-NG", action="GRIEVANCE_UPDATED").count() == 1
    db.close()

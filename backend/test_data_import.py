import os
import openpyxl
import pytest

from data_import.timetable_parser import parse_slot_time, parse_timetable_sheet
from data_import.excel_loader import load_students, load_faculty

def test_parse_slot_time_dot_format_and_24h():
    # Test dot in 10.30-11.30
    st, et = parse_slot_time("10.30-11.30")
    assert st == "10:30"
    assert et == "11:30"

    # Test morning 8:30-9:30
    st, et = parse_slot_time("8:30-9:30")
    assert st == "08:30"
    assert et == "09:30"

    # Test afternoon 1:30-2:30 -> 13:30-14:30 (24h conversion)
    st, et = parse_slot_time("1:30-2:30")
    assert st == "13:30"
    assert et == "14:30"

    # Test afternoon 2:30-3:30 -> 14:30-15:30
    st, et = parse_slot_time("2:30-3:30")
    assert st == "14:30"
    assert et == "15:30"


def test_cell_grammar_theory():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Timetable"
    # Set day header
    ws.cell(1, 5, "MONDAY")
    ws.cell(3, 5, "TIME/\nCLASS")
    ws.cell(3, 6, "D12A")
    ws.cell(4, 5, "8:30-9:30")
    ws.cell(4, 6, "TCS\n(SM)\n314")

    entries, report = parse_timetable_sheet(ws)
    assert len(entries) == 1
    e = entries[0]
    assert e.course_code == "TCS"
    assert e.faculty_initials == "SM"
    assert e.classroom == "314"
    assert e.session_kind == "Theory"
    assert e.batch is None
    assert e.start_time == "08:30"
    assert e.end_time == "09:30"
    assert e.class_id == "D12A"


def test_cell_grammar_practical():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Timetable"
    ws.cell(1, 5, "MONDAY")
    ws.cell(3, 5, "TIME/\nCLASS")
    ws.cell(3, 6, "D12A")
    ws.cell(4, 5, "10.30-11.30")
    ws.cell(4, 6, "SEA/ A/ 302/ MM\nDE /B/ 308/ ID\nCCL /C/ 310B/ PAG")

    entries, report = parse_timetable_sheet(ws)
    assert len(entries) == 3
    batches = {e.batch: e for e in entries}
    assert set(batches.keys()) == {"A", "B", "C"}
    assert batches["A"].course_code == "SEA"
    assert batches["A"].classroom == "302"
    assert batches["A"].faculty_initials == "MM"
    assert batches["A"].session_kind == "Practical"

    assert batches["B"].course_code == "DE"
    assert batches["B"].classroom == "308"
    assert batches["B"].faculty_initials == "ID"

    assert batches["C"].course_code == "CCL"
    assert batches["C"].classroom == "310B"
    assert batches["C"].faculty_initials == "PAG"


def test_cell_grammar_project():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Timetable"
    ws.cell(1, 5, "MONDAY")
    ws.cell(3, 5, "TIME/\nCLASS")
    ws.cell(3, 6, "D12A")
    ws.cell(4, 5, "1:30-2:30")
    ws.cell(4, 6, "PROJECT")

    ws.cell(3, 7, "D12B")
    ws.cell(4, 7, "A/ PROJECT")

    ws.cell(3, 8, "D12C")
    ws.cell(4, 8, "B/PROJECT")

    entries, report = parse_timetable_sheet(ws)
    assert len(entries) == 3
    by_class = {e.class_id: e for e in entries}
    assert by_class["D12A"].session_kind == "Project"
    assert by_class["D12A"].batch is None
    assert by_class["D12A"].classroom is None

    assert by_class["D12B"].session_kind == "Project"
    assert by_class["D12B"].batch == "A"

    assert by_class["D12C"].session_kind == "Project"
    assert by_class["D12C"].batch == "B"


def test_cell_grammar_joint_and_fanout():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Timetable"
    ws.cell(1, 5, "MONDAY")
    ws.cell(3, 5, "TIME/\nCLASS")
    ws.cell(3, 6, "D12A")
    ws.cell(3, 7, "D12B")
    ws.cell(3, 8, "D12C")
    ws.cell(4, 5, "9:30-10:30")
    ws.merge_cells("F4:H4")
    ws.cell(4, 6, "SD / (PRL) / 315 (D12A+D12C)")

    entries, report = parse_timetable_sheet(ws)
    assert len(entries) == 2
    classes = {e.class_id for e in entries}
    assert classes == {"D12A", "D12C"}
    for e in entries:
        assert e.course_code == "SD"
        assert e.faculty_initials == "PRL"
        assert e.classroom == "315"
        assert e.group_id is not None
        assert e.start_time == "09:30"
        assert e.end_time == "10:30"
    # Ensure both entries share the same group_id
    assert entries[0].group_id == entries[1].group_id


def test_typo_fix_d17b_to_d12b():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Timetable"
    ws.cell(1, 5, "WEDNESDAY")
    ws.cell(3, 5, "TIME/\nCLASS")
    ws.cell(3, 6, "D12A")
    ws.cell(3, 7, "D12B")
    ws.cell(3, 8, "D12C")
    ws.cell(4, 5, "8:30-9:30")
    ws.merge_cells("F4:H4")
    ws.cell(4, 6, "SD / (RJO) / 314 (D17B+D17C)")

    entries, report = parse_timetable_sheet(ws)
    assert len(entries) == 2
    classes = {e.class_id for e in entries}
    assert classes == {"D12B", "D12C"}
    assert any("normalised typo" in n and "D17B" in n for n in report.normalisations_applied)


def test_merged_cell_expansion_2hr_lab():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Timetable"
    ws.cell(1, 5, "MONDAY")
    ws.cell(3, 5, "TIME/\nCLASS")
    ws.cell(3, 6, "D12A")
    ws.cell(4, 5, "10.30-11.30")
    ws.cell(5, 5, "11:30-12:30")
    ws.merge_cells("F4:F5")
    ws.cell(4, 6, "SEA/ A/ 302/ MM")

    entries, report = parse_timetable_sheet(ws)
    assert len(entries) == 1
    e = entries[0]
    assert e.start_time == "10:30"
    assert e.end_time == "12:30"
    assert e.time == "10:30–12:30"
    assert e.source_ref == "Timetable!F4:F5"


def test_golden_real_workbook():
    xlsx_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data.xlsx")
    assert os.path.exists(xlsx_path), "data.xlsx must exist"

    wb = openpyxl.load_workbook(xlsx_path)

    # 1. Students check
    students, s_stats = load_students(wb)
    assert s_stats["total"] == 208
    assert s_stats["counts"]["D12A"] == 73
    assert s_stats["counts"]["D12B"] == 67
    assert s_stats["counts"]["D12C"] == 68

    # 2. Faculty check
    faculty, f_stats = load_faculty(wb)
    assert f_stats["total"] == 31
    known_initials = {f.initials for f in faculty}
    assert len(known_initials) == 31

    # 3. Timetable check
    entries, report = parse_timetable_sheet(wb, known_faculty_initials=known_initials)

    # Assert no entry in lunch hour (12:30-13:30)
    for e in entries:
        assert e.start_time != "12:30", f"Entry {e} found in lunch slot!"

    # Assert no unknown faculty initial
    assert len(report.unknown_faculty_initials) == 0, f"Unknown initials found: {report.unknown_faculty_initials}"

    # Assert 3 classes only
    classes = {e.class_id for e in entries}
    assert classes == {"D12A", "D12B", "D12C"}

    # Assert total entries
    assert report.total_entries == 118
